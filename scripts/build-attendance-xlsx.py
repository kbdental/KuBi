"""
The attendance sheet the clinic captures into, and KuBi reads.

The owner: "attendance will be captured and be picked from a separate Google
Sheet." So this is that sheet — three tabs whose columns are the import
contract, plus the eight controls written where somebody filling it in will
see them.

    python3 scripts/build-attendance-xlsx.py docs/clinic/kubi-attendance.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else 'docs/clinic/kubi-attendance.xlsx'
F = 'Arial'
INK = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True)
H1 = Font(name=F, size=16, bold=True); H2 = Font(name=F, size=12, bold=True)
MUTED = Font(name=F, size=9, color='666666')
HEADF = Font(name=F, size=10, bold=True, color='FFFFFF')
HEAD = PatternFill('solid', fgColor='0C5F4E')
KEY = PatternFill('solid', fgColor='FFD400'); NOTE = PatternFill('solid', fgColor='F2F4F2')
STOP = PatternFill('solid', fgColor='F8D0CC')
THIN = Side(style='thin', color='D5DBD7'); BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical='top', wrap_text=True); CTR = Alignment(horizontal='center', vertical='top')
wb = Workbook()

def head(ws, row, cols, key=()):
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(name=F, size=10, bold=True) if get_column_letter(i) in key else HEADF
        cell.fill = KEY if get_column_letter(i) in key else HEAD
        cell.border = BOX; cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 30

def widths(ws, spec):
    for col, w in spec.items(): ws.column_dimensions[col].width = w

def blanks(ws, ncols, start, n=200):
    for r in range(start, start + n):
        for i in range(1, ncols + 1):
            cell = ws.cell(row=r, column=i); cell.border, cell.font = BOX, INK

# ── Read me ──────────────────────────────────────────────────────────────
ws = wb.active; ws.title = 'Read me'
widths(ws, {'A': 3, 'B': 12, 'C': 30, 'D': 34, 'E': 60})
ws['B2'] = 'KuBi — attendance, leave and staff availability'; ws['B2'].font = H1
ws['B3'] = 'The sheet the clinic captures into. KuBi reads these three tabs and nothing else.'
ws['B3'].font = MUTED

ws['B5'] = 'The eight controls, and what KuBi does with each'; ws['B5'].font = H2
head(ws, 6, ['', 'ID', 'Activity', 'Standard', 'What KuBi does'])
CONTROLS = [
    ('ATT-001', 'Report for duty', 'Present by 9:45 AM',
     'Reads in_time. After 09:45 is late, with the minutes counted. 09:45 exactly is on time.'),
    ('ATT-002', 'Record late arrival', 'Reason entered',
     'A late arrival with a blank reason is LATE_UNEXPLAINED and escalates. Whitespace is not a reason.'),
    ('ATT-003', 'Record absence', 'Every absence classified',
     'A person MISSING from the sheet is UNKNOWN, not absent — nobody has said anything about them. '
     'That is the state this control exists to catch, and it is reported louder than a recorded absence.'),
    ('ATT-004', 'Submit planned leave', '≥7 days prior',
     'Compares from_date against requested_on. Anything shorter is flagged, including leave requested after it started.'),
    ('ATT-005', 'Approve/reject leave', 'Before the leave date',
     'A request still REQUESTED with the date inside 7 days is raised as pending.'),
    ('ATT-006', 'Apply leave rule', 'Sunday per clinic policy',
     'NOT BUILT. KuBi does not hold the policy and three common versions give different balances, '
     'so the case is detected and the calculation refused. Your ruling is needed.'),
    ('ATT-007', 'Arrange replacement', 'Critical role must have coverage',
     'Approved leave on a critical role with a blank cover_employee_code is raised. A REQUEST is not '
     'chased for cover — there is nothing to cover until it is approved.'),
    ('ATT-008', 'Daily manpower check', 'Required positions covered',
     'Runs at opening, on the Clinic readiness screen. Names what STOPS when a position is empty, '
     'not just that it is empty. Somebody late still counts as being in the building.'),
]
r = 7
for cid, act, std, does in CONTROLS:
    for i, v in enumerate([None, cid, act, std, does], start=1):
        cell = ws.cell(row=r, column=i, value=v)
        cell.font, cell.border, cell.alignment = INK, BOX, WRAP
    ws.cell(row=r, column=2).font = BOLD
    if cid == 'ATT-006': ws.cell(row=r, column=2).fill = STOP
    ws.row_dimensions[r].height = 40
    r += 1

r += 1
ws.cell(row=r, column=2, value='Two rules for whoever fills this in').font = H2
r += 1
for line in [
    'LEAVE A CELL BLANK RATHER THAN GUESSING. A blank in_time is "we do not know", and KuBi says so. '
    'Typing 00:00 or 09:00 to tidy it up creates a person who was here when they were not.',
    'A LATE ARRIVAL NEEDS A REASON IN THE SAME ROW. Not in a message, not verbally. The reason column '
    'is what stops ATT-002 escalating, and it is the only place KuBi looks.',
]:
    cell = ws.cell(row=r, column=2, value=line)
    cell.font, cell.alignment = INK, WRAP
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 34
    r += 1

r += 1
ws.cell(row=r, column=2, value='One thing your own standards cannot both do').font = H2
r += 1
cell = ws.cell(row=r, column=2, value=(
    'ATT-001 says present by 09:45. The opening procedure says the morning takes 50 minutes and the '
    'sterilisation cycle 75. With a 10:00 first patient the technician had to start at 08:45 — an hour '
    'before attendance asks them to be here, and housekeeping, both assistants and reception are all '
    '35 minutes short. KuBi computes this rather than describing it, and shows it on the readiness '
    'screen. Nothing overrules ATT-001; it needs your decision on which standard moves.'))
cell.font, cell.alignment = INK, WRAP
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
ws.row_dimensions[r].height = 76

# ── Staff ────────────────────────────────────────────────────────────────
ws = wb.create_sheet('Staff')
widths(ws, {'A': 18, 'B': 24, 'C': 30, 'D': 12})
ws['A1'] = 'Staff'; ws['A1'].font = H1
ws['A2'] = 'One row per person. Roles decide what KuBi expects of them; names never appear in the rules.'
ws['A2'].font = MUTED
head(ws, 4, ['employee_code', 'name', 'roles', 'active'])
ex = ['KB-014', 'A. Sharma', 'DENTAL_ASSISTANT', 'YES']
for i, v in enumerate(ex, start=1):
    c = ws.cell(row=5, column=i, value=v); c.font = Font(name=F, size=10, italic=True, color='777777'); c.border = BOX
ws.cell(row=4, column=3).comment = Comment(
    'One or more, separated by a comma: TREATING_DOCTOR, CLINIC_MANAGER, RECEPTION, '
    'SENIOR_ASSISTANT, DENTAL_ASSISTANT, STERILIZATION_TECHNICIAN, HOUSEKEEPING, '
    'LAB_COORDINATOR, INVENTORY_COORDINATOR.', 'KuBi', width=320, height=120)
blanks(ws, 4, 6); ws.freeze_panes = 'A5'; ws.auto_filter.ref = 'A4:D205'

# ── Attendance ───────────────────────────────────────────────────────────
ws = wb.create_sheet('Attendance')
widths(ws, {'A': 18, 'B': 14, 'C': 12, 'D': 12, 'E': 14, 'F': 46})
ws['A1'] = 'Attendance'; ws['A1'].font = H1
ws['A2'] = 'One row per person per day. A person with no row today is UNKNOWN, which is not the same as absent.'
ws['A2'].font = MUTED
head(ws, 4, ['employee_code', 'date', 'in_time', 'out_time', 'status', 'reason'], key=('C', 'F'))
ex = ['KB-014', '2026-08-11', '09:52', '18:40', 'LATE', 'Train delayed at Andheri']
for i, v in enumerate(ex, start=1):
    c = ws.cell(row=5, column=i, value=v); c.font = Font(name=F, size=10, italic=True, color='777777'); c.border = BOX
ws.cell(row=4, column=3).comment = Comment(
    '24-hour or am/pm both read: 09:52, 9:52, 9.52, 9:52 am, 1:05 pm. '
    'BLANK MEANS NOT KNOWN — do not type 00:00.', 'KuBi', width=300, height=100)
ws.cell(row=4, column=5).comment = Comment(
    'PRESENT · LATE · ABSENT · LEAVE · HOLIDAY. Leave blank only if the row is being filled in later; '
    'a blank row with a blank status reads as UNKNOWN.', 'KuBi', width=320, height=110)
ws.cell(row=4, column=6).comment = Comment(
    'ATT-002. Required whenever in_time is after 09:45. A blank here escalates.',
    'KuBi', width=300, height=80)
blanks(ws, 6, 6); ws.freeze_panes = 'B5'; ws.auto_filter.ref = 'A4:F205'

# ── Leave ────────────────────────────────────────────────────────────────
ws = wb.create_sheet('Leave')
widths(ws, {'A': 18, 'B': 13, 'C': 13, 'D': 16, 'E': 14, 'F': 14, 'G': 22, 'H': 14})
ws['A1'] = 'Leave'; ws['A1'].font = H1
ws['A2'] = 'One row per request. KuBi reads notice, decision and cover from it.'
ws['A2'].font = MUTED
head(ws, 4, ['employee_code', 'from_date', 'to_date', 'kind', 'status',
             'requested_on', 'cover_employee_code', 'spans_sunday'], key=('F', 'G'))
ex = ['KB-014', '2026-08-20', '2026-08-22', 'ANNUAL', 'APPROVED', '2026-08-05', 'KB-021', 'NO']
for i, v in enumerate(ex, start=1):
    c = ws.cell(row=5, column=i, value=v); c.font = Font(name=F, size=10, italic=True, color='777777'); c.border = BOX
ws.cell(row=4, column=6).comment = Comment(
    'ATT-004. The day the request was MADE, not the day it starts. Without it the seven-day '
    'notice rule cannot be checked at all.', 'KuBi', width=320, height=100)
ws.cell(row=4, column=7).comment = Comment(
    'ATT-007. Who is covering. Blank on approved leave for a critical role is raised as a gap.',
    'KuBi', width=300, height=90)
ws.cell(row=4, column=8).comment = Comment(
    'ATT-006. YES where the leave spans a Saturday and the following Monday. KuBi detects the case '
    'and REFUSES to compute the balance until the clinic policy is stated.', 'KuBi', width=330, height=110)
blanks(ws, 8, 6); ws.freeze_panes = 'B5'; ws.auto_filter.ref = 'A4:H205'

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(f'{OUT} — {len(wb.worksheets)} sheets')
